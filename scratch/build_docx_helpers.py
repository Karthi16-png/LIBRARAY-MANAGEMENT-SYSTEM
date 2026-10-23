"""
Script to generate the comprehensive, professional Word Document (.docx) report:
LIBRARY MANAGEMENT SYSTEM - DIGITAL ASSIGNMENT 2
Authors: KARTHIK VALLIAPPAN C (25BAI1051), NITHIN RANGASWAMY (25BAI1210), NATRAJ (25BAI1079)
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    """Sets the background color of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets internal padding for a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tc_pr.append(tc_mar)

def set_table_borders(table, color="D1D5DB", sz="4", val="single"):
    """Applies clean borders to a table."""
    tbl_pr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tbl_pr.append(borders)

def add_code_block(doc, code_text):
    """Adds a formatted code block with monospace font and shaded container."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    # Border for code block
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:color="CBD5E1"/>
            <w:left w:val="single" w:sz="18" w:color="2563EB"/>
            <w:bottom w:val="single" w:sz="6" w:color="CBD5E1"/>
            <w:right w:val="single" w:sz="6" w:color="CBD5E1"/>
        </w:tcBorders>
    ''')
    tc_pr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(code_text)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_query_box(doc, q_num, title, sql, result_text, purpose):
    """Adds a standard formatted query presentation card."""
    h = doc.add_heading(f"{q_num}. {title}", level=3)
    h.paragraph_format.space_before = Pt(10)
    h.paragraph_format.space_after = Pt(4)
    h.runs[0].font.color.rgb = RGBColor(27, 54, 93)
    
    # SQL Query Block
    add_code_block(doc, sql)
    
    # Query Result Block
    p_res_title = doc.add_paragraph()
    p_res_title.paragraph_format.space_before = Pt(2)
    p_res_title.paragraph_format.space_after = Pt(2)
    r_lbl = p_res_title.add_run("Query Result:")
    r_lbl.bold = True
    r_lbl.font.size = Pt(9.5)
    r_lbl.font.color.rgb = RGBColor(51, 65, 85)
    
    add_code_block(doc, result_text)
    
    # Purpose
    p_purp = doc.add_paragraph()
    p_purp.paragraph_format.space_before = Pt(2)
    p_purp.paragraph_format.space_after = Pt(8)
    r_purp_lbl = p_purp.add_run("Purpose: ")
    r_purp_lbl.bold = True
    r_purp_lbl.font.size = Pt(9.5)
    r_purp_lbl.font.color.rgb = RGBColor(30, 64, 175)
    r_purp_txt = p_purp.add_run(purpose)
    r_purp_txt.font.size = Pt(9.5)
    r_purp_txt.font.color.rgb = RGBColor(55, 65, 81)

print("Helper functions defined successfully.")
